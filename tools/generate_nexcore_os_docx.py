from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SOURCE = Path(__file__).resolve().parents[1] / "Nexcore OS" / "ARCHITECTURE.md"
OUTPUT = Path.home() / "Downloads" / "NexCore OS Architecture.docx"


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_inline(paragraph, text, font_name=None):
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(30, 64, 175)
        elif part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        else:
            run = paragraph.add_run(part)
            if font_name:
                run.font.name = font_name


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(5)

    colors = {
        "Title": RGBColor(15, 23, 42),
        "Heading 1": RGBColor(15, 76, 129),
        "Heading 2": RGBColor(8, 116, 122),
        "Heading 3": RGBColor(51, 65, 85),
    }
    for name, color in colors.items():
        style = doc.styles[name]
        style.font.name = "Aptos Display"
        style.font.color.rgb = color


def add_cover(doc):
    doc.add_paragraph()
    doc.add_paragraph()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("NexCore OS")
    run.bold = True
    run.font.name = "Aptos Display"
    run.font.size = Pt(34)
    run.font.color.rgb = RGBColor(15, 76, 129)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Execution Intelligence Operating System\nArchitecture Specification")
    run.font.name = "Aptos Display"
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(71, 85, 105)

    doc.add_paragraph()
    summary = doc.add_paragraph()
    summary.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = summary.add_run(
        "A practical platform architecture for designing, executing, observing, "
        "diagnosing, and improving automated work across web, API, mobile, desktop, and data systems."
    )
    run.italic = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(71, 85, 105)

    doc.add_paragraph()
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run("Architecture baseline • Version 1.0 • July 2026").font.color.rgb = RGBColor(100, 116, 139)
    doc.add_page_break()


def render_markdown(doc, markdown):
    lines = markdown.splitlines()
    i = 0
    in_code = False
    code_lines = []

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            if not in_code:
                in_code = True
                code_lines = []
            else:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.25)
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(8)
                run = p.add_run("\n".join(code_lines))
                run.font.name = "Consolas"
                run.font.size = Pt(8.5)
                run.font.color.rgb = RGBColor(30, 41, 59)
                in_code = False
            i += 1
            continue

        if in_code:
            code_lines.append(line)
            i += 1
            continue

        if stripped.startswith("|") and i + 1 < len(lines) and re.match(r"^\s*\|?\s*:?-+", lines[i + 1]):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append([cell.strip() for cell in lines[i].strip().strip("|").split("|")])
                i += 1
            if len(rows) >= 2:
                data = [rows[0]] + rows[2:]
                table = doc.add_table(rows=len(data), cols=max(len(r) for r in data))
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                table.style = "Table Grid"
                for r_idx, row in enumerate(data):
                    for c_idx, value in enumerate(row):
                        cell = table.cell(r_idx, c_idx)
                        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                        p = cell.paragraphs[0]
                        add_inline(p, value)
                        if r_idx == 0:
                            shade_cell(cell, "0F4C81")
                            for run in p.runs:
                                run.bold = True
                                run.font.color.rgb = RGBColor(255, 255, 255)
                doc.add_paragraph()
            continue

        heading = re.match(r"^(#{1,6})\s+(.+)$", stripped)
        if heading:
            level = min(len(heading.group(1)), 3)
            text = heading.group(2)
            if level == 1 and text == "NexCore OS Architecture":
                i += 1
                continue
            doc.add_heading(text, level=level)
        elif re.match(r"^[-*]\s+", stripped):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, re.sub(r"^[-*]\s+", "", stripped))
        elif re.match(r"^\d+\.\s+", stripped):
            p = doc.add_paragraph(style="List Number")
            add_inline(p, re.sub(r"^\d+\.\s+", "", stripped))
        elif stripped:
            p = doc.add_paragraph()
            add_inline(p, stripped)
        i += 1


def main():
    markdown = SOURCE.read_text(encoding="utf-8")
    doc = Document()
    configure_document(doc)
    add_cover(doc)
    render_markdown(doc, markdown)

    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("NexCore OS Architecture • Confidential working specification")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(100, 116, 139)

    doc.core_properties.title = "NexCore OS Architecture"
    doc.core_properties.subject = "Execution Intelligence Operating System Architecture"
    doc.core_properties.author = "NexCore"
    doc.core_properties.keywords = "NexCore, architecture, execution intelligence, operating system"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
