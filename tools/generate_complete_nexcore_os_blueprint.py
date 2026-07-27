from pathlib import Path
import textwrap

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path.home() / "Downloads"
ASSETS = Path(__file__).resolve().parents[1] / ".tmp-nexcore-blueprints"
DOCX = OUT / "NexCore OS Complete Orchestration Blueprint.docx"
ASSETS.mkdir(exist_ok=True)

W, H = 2200, 1280
BG, GRID, PANEL = "#06111f", "#10304d", "#0c2137"
WHITE, MUTED = "#eef7ff", "#9bb4ca"
BLUE, CYAN, GREEN = "#3787ff", "#38d9ff", "#42e6a4"
AMBER, PURPLE, RED = "#ffc857", "#b68cff", "#ff6b7a"


def fnt(size, bold=False):
    path = Path("C:/Windows/Fonts/seguisb.ttf" if bold else "C:/Windows/Fonts/segoeui.ttf")
    return ImageFont.truetype(str(path), size) if path.exists() else ImageFont.load_default()


FT, FS, FH, FB, FM, FX = fnt(48, True), fnt(21), fnt(25, True), fnt(20, True), fnt(17), fnt(14)


def wrap(text, width=26):
    return "\n".join(textwrap.wrap(text, width=width))


def base(title, subtitle, index):
    img = Image.new("RGB", (W, H), BG)
    d = ImageDraw.Draw(img)
    for x in range(0, W, 50): d.line((x, 0, x, H), fill=GRID)
    for y in range(0, H, 50): d.line((0, y, W, y), fill=GRID)
    d.text((70, 45), title, font=FT, fill=WHITE)
    d.text((72, 108), subtitle, font=FS, fill=CYAN)
    d.text((1960, 58), f"BP-{index:02d}", font=FH, fill=MUTED)
    return img, d


def box(d, xy, title, body="", color=BLUE, fill=PANEL):
    d.rounded_rectangle(xy, 18, fill=fill, outline=color, width=3)
    x1, y1, x2, y2 = xy
    d.line((x1+14, y1+8, x2-14, y1+8), fill=color, width=5)
    tb = d.textbbox((0, 0), title, font=FB)
    d.text(((x1+x2-(tb[2]-tb[0]))/2, y1+22), title, font=FB, fill=WHITE)
    yy = y1+56
    for line in body.split("\n"):
        bb = d.textbbox((0, 0), line, font=FM)
        d.text(((x1+x2-(bb[2]-bb[0]))/2, yy), line, font=FM, fill=MUTED)
        yy += 24


def lane(d, xy, title, color):
    d.rounded_rectangle(xy, 18, fill="#071827", outline=color, width=2)
    d.text((xy[0]+18, xy[1]+12), title, font=FH, fill=color)


def arrow(d, a, b, color=CYAN, label="", dashed=False):
    import math
    x1, y1 = a; x2, y2 = b
    if dashed:
        for i in range(0, 20, 2):
            aa, bb = i/20, (i+1)/20
            d.line((x1+(x2-x1)*aa, y1+(y2-y1)*aa, x1+(x2-x1)*bb, y1+(y2-y1)*bb), fill=color, width=4)
    else: d.line((x1, y1, x2, y2), fill=color, width=4)
    ang, size = math.atan2(y2-y1, x2-x1), 14
    d.polygon([(x2,y2),(x2-size*math.cos(ang-.55),y2-size*math.sin(ang-.55)),(x2-size*math.cos(ang+.55),y2-size*math.sin(ang+.55))], fill=color)
    if label:
        mx, my = (x1+x2)//2, (y1+y2)//2
        bb = d.textbbox((0,0), label, font=FX)
        d.rounded_rectangle((mx-(bb[2]-bb[0])/2-6,my-18,mx+(bb[2]-bb[0])/2+6,my+3),5,fill=BG)
        d.text((mx-(bb[2]-bb[0])/2,my-16),label,font=FX,fill=color)


def footer(d, text):
    d.text((70, 1230), text, font=FM, fill=MUTED)


def save(img, n, slug):
    path = ASSETS / f"{n:02d}-{slug}.png"
    img.save(path, quality=95)
    return path


def diagrams():
    paths=[]
    # 1 system context
    img,d=base("NEXCORE OS — SYSTEM CONTEXT","Actors, platform boundary, execution targets and external systems",1)
    lane(d,(55,175,2145,1175),"NEXCORE OS PLATFORM BOUNDARY",CYAN)
    for xy,t,b,c in [((95,310,380,445),"Teams","QA • Dev • Ops",BLUE),((95,540,380,675),"CI/CD","GitHub • Jenkins",AMBER),((730,260,1450,480),"Experience + Control Plane","Workspace • APIs • Identity • Catalog\nPolicy • Audit • Realtime",CYAN),((730,610,1450,850),"Orchestration Kernel","Temporal workflows • Scheduler\nLifecycle • Retry • Approval • Replay",PURPLE),((1700,280,2070,430),"Enterprise Systems","Jira • Git • Chat • BI",AMBER),((1650,650,2070,850),"Execution Targets","Web • API • Mobile\nDesktop • Data",GREEN),((760,960,1420,1110),"Evidence + Intelligence","Artifacts • Observability • Diagnosis • Repair",PURPLE)]: box(d,xy,t,b,c)
    arrow(d,(380,375),(730,350),BLUE,"author/run"); arrow(d,(380,607),(730,420),AMBER,"trigger")
    arrow(d,(1450,350),(1700,350),AMBER,"integrate"); arrow(d,(1090,480),(1090,610),CYAN,"commands")
    arrow(d,(1450,730),(1650,750),GREEN,"execute"); arrow(d,(1860,850),(1420,1015),PURPLE,"evidence")
    arrow(d,(1090,960),(1090,850),PURPLE,"diagnose/replay")
    footer(d,"Boundary rule: NexCore OS coordinates operating systems and runtimes; it does not replace Windows, Linux, browsers, or device platforms.")
    paths.append(save(img,1,"system-context"))

    # 2 service/container architecture
    img,d=base("CONTAINER AND SERVICE ARCHITECTURE","Deployable responsibilities and authoritative data ownership",2)
    lane(d,(55,170,2145,390),"EDGE / EXPERIENCE",BLUE); lane(d,(55,415,2145,715),"CONTROL / ORCHESTRATION",CYAN); lane(d,(55,740,2145,995),"EXECUTION / INTELLIGENCE",GREEN); lane(d,(55,1020,2145,1185),"DATA",PURPLE)
    specs=[((105,235,430,350),"Next.js UI","Workspace + React Flow",BLUE),((550,235,875,350),"API Gateway","REST + realtime",BLUE),((995,235,1320,350),"Identity","Keycloak / OIDC",AMBER),((1440,235,2050,350),"External API + SDK","Versioned OpenAPI contracts",BLUE),((105,500,480,650),"Control Plane","NestJS modular monolith\nCatalog • Policy • Audit",CYAN),((590,500,965,650),"Temporal Kernel","Durable workflow ownership",PURPLE),((1075,500,1450,650),"Scheduler","Capabilities + leases",GREEN),((1560,500,2050,650),"NATS Fabric","Outbox • projections • integration",AMBER),((105,815,500,940),"Runtime Agents","Web • API • Mobile • Desktop",GREEN),((650,815,1045,940),"Python Workers","Plugins + automation",GREEN),((1195,815,1590,940),"AI Workers","Diagnosis + proposals",PURPLE),((1740,815,2050,940),"Collectors","OTel + evidence",BLUE),((105,1070,420,1150),"PostgreSQL","Canonical",BLUE),((520,1070,835,1150),"Temporal DB","Durability",PURPLE),((935,1070,1250,1150),"MinIO","Artifacts",GREEN),((1350,1070,1665,1150),"Qdrant","Derived",PURPLE),((1765,1070,2080,1150),"Secrets","References",AMBER)]
    for x,t,b,c in specs: box(d,x,t,b,c)
    arrow(d,(430,292),(550,292),BLUE); arrow(d,(292,500),(292,350),CYAN); arrow(d,(480,575),(590,575),CYAN); arrow(d,(965,575),(1075,575),GREEN); arrow(d,(1450,575),(1560,575),AMBER)
    arrow(d,(1260,650),(350,815),GREEN,"lease"); arrow(d,(1820,650),(1390,815),PURPLE,"events"); arrow(d,(300,940),(260,1070),BLUE); arrow(d,(850,940),(1090,1070),GREEN); arrow(d,(1390,940),(1510,1070),PURPLE)
    footer(d,"Start as a modular control plane; split services only when scaling, security isolation, or team ownership provides measurable benefit.")
    paths.append(save(img,2,"containers"))

    # 3 orchestration sequence
    img,d=base("END-TO-END ORCHESTRATION SEQUENCE","Command acceptance through durable completion and realtime projection",3)
    actors=["User / CI","API","Policy","Temporal","Scheduler","Agent","Storage","AI / UI"]
    xs=[120,390,660,930,1200,1470,1740,2010]
    for x,a in zip(xs,actors):
        d.text((x-55,185),a,font=FB,fill=CYAN); d.line((x,225,x,1160),fill="#31516b",width=2)
    events=[(0,1,280,"Start + idempotency key",BLUE),(1,2,350,"Authorize + quota",AMBER),(2,3,420,"Create durable workflow",CYAN),(3,4,500,"Request capability",GREEN),(4,5,575,"Issue bounded lease",GREEN),(5,6,655,"Upload evidence",PURPLE),(5,3,735,"Structured result",CYAN),(3,6,815,"Persist transition",BLUE),(3,7,895,"Publish projection",PURPLE),(7,3,990,"Approval / signal",AMBER),(3,5,1070,"Retry or next node",GREEN)]
    for a,b,y,l,c in events: arrow(d,(xs[a],y),(xs[b],y),c,l)
    footer(d,"All command handlers and callbacks are idempotent. Database state plus event publication uses an outbox boundary.")
    paths.append(save(img,3,"sequence"))

    # 4 temporal kernel
    img,d=base("TEMPORAL ORCHESTRATION KERNEL","Deterministic workflow ownership, activity boundaries, signals and queries",4)
    lane(d,(55,175,2145,1145),"DURABLE KERNEL",PURPLE)
    specs=[((105,260,490,410),"Execution Workflow","Canonical state machine\nDAG readiness + terminal state",PURPLE),((640,240,1030,430),"Node Coordinator","Attempt • retry • timeout\nchild workflow / activity",CYAN),((1180,240,1570,430),"Signals","Cancel • pause • approve\nresume • repair",AMBER),((1720,240,2080,430),"Queries","Status • graph • progress",BLUE),((105,610,490,790),"Deterministic Logic","No network • no wall-clock\nversioned workflow changes",RED),((640,610,1030,790),"Activities","DB • NATS • scheduler\nplugin dispatch • artifacts",GREEN),((1180,610,1570,790),"Failure Policy","Classify • backoff • jitter\ncompensate • dead letter",AMBER),((1720,610,2080,790),"Continue-As-New","Bound history for long runs",CYAN),((520,940,940,1080),"Search Attributes","Tenant-safe operational lookup",BLUE),((1110,940,1530,1080),"Workflow Versioning","Patch markers + replay tests",PURPLE)]
    for x,t,b,c in specs: box(d,x,t,b,c)
    arrow(d,(490,335),(640,335),CYAN); arrow(d,(1030,335),(1180,335),AMBER,dashed=True); arrow(d,(1570,335),(1720,335),BLUE,dashed=True)
    arrow(d,(835,430),(835,610),GREEN,"activity call"); arrow(d,(1375,430),(1375,610),AMBER,"policy"); arrow(d,(300,410),(300,610),RED,"invariant")
    footer(d,"Temporal owns transitions; activities perform external I/O. Workflow changes require deterministic replay compatibility tests.")
    paths.append(save(img,4,"temporal-kernel"))

    # 5 state machines
    img,d=base("EXECUTION AND NODE STATE MACHINES","Allowed lifecycle transitions, terminal ownership and retry semantics",5)
    lane(d,(55,175,1070,1160),"EXECUTION LIFECYCLE",CYAN); lane(d,(1130,175,2145,1160),"NODE ATTEMPT LIFECYCLE",GREEN)
    left=[((115,300,390,390),"VALIDATING","schema • policy",BLUE),((500,300,775,390),"QUEUED","accepted",CYAN),((500,475,775,565),"PROVISIONING","allocate runtime",AMBER),((500,650,775,740),"RUNNING","traverse DAG",GREEN),((130,830,405,920),"AWAITING APPROVAL","durable wait",AMBER),((500,830,775,920),"PAUSED","signal wait",PURPLE),((800,830,1020,920),"CANCELLING","propagate",RED),((115,1010,335,1095),"FAILED","terminal",RED),((420,1010,640,1095),"SUCCEEDED","terminal",GREEN),((725,1010,945,1095),"CANCELLED","terminal",MUTED)]
    right=[((1190,300,1450,390),"CREATED","materialized",BLUE),((1550,300,1810,390),"READY","deps satisfied",CYAN),((1550,475,1810,565),"LEASED","agent assigned",AMBER),((1550,650,1810,740),"RUNNING","heartbeat",GREEN),((1190,830,1450,920),"RETRY WAIT","bounded backoff",AMBER),((1550,830,1810,920),"COMPLETED","terminal",GREEN),((1850,830,2090,920),"CANCELLED","terminal",MUTED),((1190,1010,1450,1095),"FAILED","terminal/exhausted",RED),((1550,1010,1810,1095),"SKIPPED","branch policy",PURPLE)]
    for x,t,b,c in left+right: box(d,x,t,b,c)
    for a,b in [((390,345),(500,345)),((637,390),(637,475)),((637,565),(637,650)),((500,695),(405,830)),((637,740),(637,830)),((775,695),(910,830)),((637,920),(530,1010)),((405,920),(225,1010)),((910,920),(835,1010))]: arrow(d,a,b,CYAN)
    for a,b in [((1450,345),(1550,345)),((1680,390),(1680,475)),((1680,565),(1680,650)),((1550,695),(1450,830)),((1680,740),(1680,830)),((1810,695),(1970,830)),((1320,920),(1320,1010)),((1680,920),(1680,1010)),((1320,830),(1550,695))]: arrow(d,a,b,GREEN)
    footer(d,"Attempts are immutable records. A retry creates a new attempt; it does not overwrite the failed attempt or its evidence.")
    paths.append(save(img,5,"state-machines"))

    # 6 agent scheduling
    img,d=base("RUNTIME AGENT SCHEDULING AND RECOVERY","Capability discovery, policy filtering, leases, heartbeat and orphan recovery",6)
    specs=[((90,245,410,390),"Agent Registry","Identity • platform\ncapability versions",BLUE),((520,245,840,390),"Health + Capacity","heartbeat • slots\nlabels • locality",GREEN),((950,245,1270,390),"Policy Filter","tenant • network zone\ndata residency",AMBER),((1380,245,1700,390),"Scoring","compatibility • load\ncost • affinity",CYAN),((1810,245,2110,390),"Lease","TTL • token\nrenew • revoke",PURPLE),((190,650,570,820),"Cloud Pool","ephemeral web/API agents",GREEN),((700,650,1080,820),"Device Lab","mobile reservations",AMBER),((1210,650,1590,820),"Customer Network","desktop/internal targets",PURPLE),((1720,650,2050,820),"Recovery Controller","expire lease • classify\nreschedule or fail",RED),((690,1010,1510,1135),"Scheduler Decision Record","requested capabilities + candidates + policy exclusions + selected agent + lease history",BLUE)]
    for x,t,b,c in specs: box(d,x,t,b,c)
    for a,b in [((410,315),(520,315)),((840,315),(950,315)),((1270,315),(1380,315)),((1700,315),(1810,315))]: arrow(d,a,b,CYAN)
    for x in [380,890,1400]: arrow(d,(1960,390),(x,650),GREEN,"lease" if x==890 else "")
    for x in [380,890,1400]: arrow(d,(x,820),(1880,650),RED,"heartbeat" if x==1400 else "",True)
    arrow(d,(1880,820),(1100,1010),RED,"record recovery")
    footer(d,"A lease is not ownership. Lost heartbeat expires the lease; retry depends on operation idempotency and failure classification.")
    paths.append(save(img,6,"agent-scheduling"))

    # 7 AI repair loop
    img,d=base("EVIDENCE-BASED AI DIAGNOSIS AND REPAIR","Grounded analysis, risk classification, approval and versioned replay",7)
    stages=[((75,300,340,440),"Failure Event","typed category\ntrace context",RED),((425,300,690,440),"Evidence Bundle","logs • screenshot\nDOM • trace • history",BLUE),((775,300,1040,440),"Diagnosis","hypotheses + citations\nconfidence",PURPLE),((1125,300,1390,440),"Repair Proposal","structured patch\nimpact scope",PURPLE),((1475,300,1740,440),"Policy Gate","allow • approve\nprohibit",AMBER),((1825,300,2125,440),"New Version","immutable change\nrollback link",GREEN)]
    for x,t,b,c in stages: box(d,x,t,b,c)
    for i in range(len(stages)-1): arrow(d,(stages[i][0][2],370),(stages[i+1][0][0],370),PURPLE)
    branches=[((160,700,500,850),"Auto-apply","Only low-risk, reversible\nand explicitly permitted",GREEN),((680,700,1020,850),"Human approval","Evidence, diff, confidence\nand blast radius",AMBER),((1200,700,1540,850),"Reject / Escalate","Unsafe, unsupported or\ninsufficient evidence",RED),((1720,700,2060,850),"Replay Evaluation","Affected subgraph\ncompare outcome + cost",CYAN)]
    for x,t,b,c in branches: box(d,x,t,b,c)
    arrow(d,(1605,440),(330,700),GREEN); arrow(d,(1605,440),(850,700),AMBER); arrow(d,(1605,440),(1370,700),RED); arrow(d,(1975,440),(1890,700),CYAN)
    box(d,(550,1010,1650,1135),"Learning Record","proposal • approver • before/after evidence • result • rollback • model/prompt version",BLUE)
    for x in [330,850,1370,1890]: arrow(d,(x,850),(1100,1010),BLUE,dashed=True)
    footer(d,"AI never mutates the active workflow or execution state directly. Every recommendation is cited, policy-evaluated and auditable.")
    paths.append(save(img,7,"ai-repair"))

    # 8 security
    img,d=base("SECURITY, TENANCY AND TRUST BOUNDARIES","Identity propagation, least privilege, secret handling and plugin isolation",8)
    lane(d,(55,180,2145,1120),"ZERO-TRUST PLATFORM",AMBER)
    specs=[((100,270,420,420),"Human / Service Identity","OIDC • MFA • workload ID",BLUE),((540,270,860,420),"API Enforcement","tenant from claims\nRBAC/ABAC • quota",AMBER),((980,270,1300,420),"Policy Engine","resource + action\ncontext + risk",AMBER),((1420,270,1740,420),"Audit Ledger","append-only decisions",PURPLE),((1860,270,2090,420),"KMS / Secrets","reference only",RED),((140,640,520,820),"Tenant Data Boundary","RLS • scoped queries\nobject prefixes",BLUE),((660,640,1040,820),"Agent Trust Boundary","mTLS • short token\nlease-scoped access",GREEN),((1180,640,1560,820),"Plugin Sandbox","signed manifest\npermissions • egress limits",PURPLE),((1700,640,2050,820),"AI Boundary","redaction • tenant retrieval\nprovider controls",PURPLE),((560,950,1640,1060),"Security Evidence","authentication • authorization • secret access • plugin install • approval • export • administrative change",AMBER)]
    for x,t,b,c in specs: box(d,x,t,b,c)
    for a,b in [((420,345),(540,345)),((860,345),(980,345)),((1300,345),(1420,345)),((1740,345),(1860,345))]: arrow(d,a,b,AMBER)
    for x in [330,850,1370,1875]: arrow(d,(1140,420),(x,640),CYAN,"scoped" if x==850 else "")
    for x in [330,850,1370,1875]: arrow(d,(x,820),(1100,950),AMBER,dashed=True)
    footer(d,"Tenant identity is derived from verified claims, never request payloads. Secrets are resolved just-in-time and excluded from logs and artifacts.")
    paths.append(save(img,8,"security"))

    # 9 data
    img,d=base("DATA, EVENT AND ARTIFACT ARCHITECTURE","Sources of truth, derived stores, lineage, retention and delivery semantics",9)
    specs=[((90,260,440,430),"PostgreSQL","Canonical resources\nexecution state • audit\ntransaction + outbox",BLUE),((570,260,920,430),"NATS JetStream","Versioned integration events\nat-least-once delivery",AMBER),((1050,260,1400,430),"Projections","Realtime UI • search\nnotifications • analytics",CYAN),((1530,260,1880,430),"Consumers","Idempotent handlers\nconsumer checkpoints",GREEN),((130,690,500,870),"MinIO","Artifact bytes\nchecksum • retention\nlegal hold",GREEN),((640,690,1010,870),"Artifact Catalog","metadata • lineage\ncontent type • access",BLUE),((1150,690,1520,870),"Qdrant","Derived embeddings\nrebuildable index",PURPLE),((1660,690,2030,870),"Warehouse / BI","De-identified export\noperational analytics",CYAN),((600,1030,1600,1145),"Data Contract Governance","schema registry • compatibility • deletion policy • residency • backup/restore tests • RPO/RTO",PURPLE)]
    for x,t,b,c in specs: box(d,x,t,b,c)
    arrow(d,(440,345),(570,345),AMBER,"outbox"); arrow(d,(920,345),(1050,345),CYAN); arrow(d,(1400,345),(1530,345),GREEN)
    arrow(d,(315,430),(315,690),GREEN,"signed upload"); arrow(d,(500,780),(640,780),BLUE,"metadata"); arrow(d,(1010,780),(1150,780),PURPLE,"derive"); arrow(d,(1520,780),(1660,780),CYAN,"approved export")
    for x in [315,825,1335,1845]: arrow(d,(x,870),(1100,1030),PURPLE,dashed=True)
    footer(d,"PostgreSQL is canonical. MinIO is evidentiary. Qdrant and UI projections are derived and must be rebuildable.")
    paths.append(save(img,9,"data"))

    # 10 deployment
    img,d=base("DEPLOYMENT, SCALING AND RESILIENCE TOPOLOGY","Cloud control plane with isolated execution zones and tested recovery",10)
    lane(d,(55,175,1080,1145),"NEXCORE MANAGED CLOUD",CYAN); lane(d,(1120,175,2145,1145),"CUSTOMER / SPECIALIZED EXECUTION ZONES",GREEN)
    specs=[((100,260,450,400),"Edge","WAF • ingress • UI\nAPI + realtime replicas",BLUE),((580,260,930,400),"Control Cluster","control plane • scheduler\nTemporal/NATS clients",CYAN),((100,535,450,700),"Orchestration Cluster","Temporal frontend/history\nworkers • visibility",PURPLE),((580,535,930,700),"Data Services","PostgreSQL HA • MinIO\nNATS • Qdrant • secrets",AMBER),((100,850,450,1015),"Cloud Runtime Pools","autoscaled isolated jobs\nweb + API",GREEN),((580,850,930,1015),"Observability","OTel • metrics • logs\nalerts • SLOs",BLUE),((1180,270,1580,440),"Customer Agent Gateway","outbound mTLS connection\nno inbound firewall opening",CYAN),((1690,270,2070,440),"Internal Desktop Pool","Windows agents near AUT",PURPLE),((1180,650,1580,820),"Mobile Device Lab","reservation + health\nphysical/virtual devices",AMBER),((1690,650,2070,820),"Private Data Pool","API/data execution\nresidency constrained",GREEN),((1290,960,1960,1080),"Recovery Region","configuration + backups • restore automation • failover runbooks • chaos tests",RED)]
    for x,t,b,c in specs: box(d,x,t,b,c)
    arrow(d,(450,330),(580,330),CYAN); arrow(d,(755,400),(275,535),PURPLE); arrow(d,(755,400),(755,535),AMBER); arrow(d,(275,700),(275,850),GREEN); arrow(d,(755,700),(755,850),BLUE)
    arrow(d,(930,330),(1180,355),CYAN,"mTLS / NATS"); arrow(d,(1580,355),(1690,355),PURPLE); arrow(d,(1380,440),(1380,650),AMBER); arrow(d,(1580,735),(1690,735),GREEN)
    arrow(d,(755,700),(1625,960),RED,"backup/config",True)
    footer(d,"Execution pools scale independently. Desktop/mobile constraints are capacity-scheduled; recovery is proven through restore and failover exercises.")
    paths.append(save(img,10,"deployment"))
    return paths


SECTIONS = [
    ("1. System context", "Defines what NexCore OS owns and what remains external.", ["Treat NexCore as an execution operating system, not a literal OS.", "Support human, CI/CD and external-system actors through the same governed command model.", "Keep runtime targets outside the kernel boundary."], ["Product sprawl beyond execution intelligence", "Unclear platform boundary causing duplicated responsibility"]),
    ("2. Container and service architecture", "Maps responsibilities to deployable units without premature microservices.", ["Keep Next.js for the experience plane.", "Keep NestJS as a modular control plane until evidence justifies replacement.", "Keep Python for automation and AI workers.", "Use language-neutral contracts between planes."], ["Distributed-monolith coupling", "Multiple implementations of lifecycle rules"]),
    ("3. End-to-end orchestration sequence", "Shows authoritative ordering, durability and realtime projection.", ["Accept commands idempotently.", "Start Temporal only after validation and canonical execution creation.", "Upload artifacts directly with short-lived credentials.", "Project state after the kernel commits transitions."], ["Duplicate execution starts", "State/event inconsistency without outbox", "Lost callbacks without idempotency"]),
    ("4. Temporal kernel internals", "Separates deterministic orchestration from side-effecting activities.", ["Workflow code owns DAG readiness and terminal state.", "Activities own databases, messaging, scheduling and runtime calls.", "Signals carry commands; queries expose read-only state.", "Replay-test every workflow-code change."], ["Nondeterministic workflow changes", "Unbounded workflow history", "Retry storms"]),
    ("5. Execution and node state machines", "Defines legal transitions and immutable attempt history.", ["Only the kernel changes state.", "A retry creates a new attempt.", "Cancellation is propagated and acknowledged.", "Terminal states cannot reopen; replay creates an explicit new run or branch."], ["Conflicting status vocabularies", "Overwritten failure evidence", "Ambiguous cancellation"]),
    ("6. Agent scheduling and recovery", "Explains capability placement across cloud and customer environments.", ["Match capability version, health, capacity, locality, policy and cost.", "Use bounded renewable leases and heartbeats.", "Persist scheduler decision records.", "Reschedule only operations proven safe or idempotent."], ["Duplicate side effects after lease loss", "Stale capability advertisements", "Customer-network reachability"]),
    ("7. AI diagnosis and repair", "Constrains intelligence to evidence-backed proposals and governed mutation.", ["Require citations and confidence for diagnoses.", "Classify repair risk through policy.", "Create a new workflow version for every accepted repair.", "Evaluate replay outcome and retain rollback linkage."], ["Hallucinated repairs", "Cross-tenant retrieval leakage", "Automation bias in approval"]),
    ("8. Security and tenancy", "Defines identity, data, agent, plugin and AI trust boundaries.", ["Derive tenant identity from verified claims.", "Use defense-in-depth tenant scoping and row-level security.", "Keep secret values out of persistence, logs and artifacts.", "Sign plugins and constrain permissions and network egress."], ["Broken object-level authorization", "Secret leakage", "Over-privileged agents/plugins"]),
    ("9. Data, events and artifacts", "Classifies canonical, evidentiary and derived stores.", ["PostgreSQL is canonical.", "MinIO stores immutable evidence bytes plus checksums and retention.", "NATS provides delivery, not canonical truth.", "Qdrant and projections are rebuildable."], ["Schema incompatibility", "Unbounded artifact cost", "Deletion/residency violations"]),
    ("10. Deployment and resilience", "Places independently scalable services and constrained runtime pools.", ["Scale control, orchestration and execution separately.", "Prefer outbound mTLS from customer agents.", "Define SLOs and capacity policies per pool.", "Test restore, failover and degraded operation."], ["Single-region dependency", "Desktop/mobile capacity bottlenecks", "Untested backups"]),
]


def shade(cell, color):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement("w:shd"); shd.set(qn("w:fill"),color); tcPr.append(shd)


def add_table(doc, headers, rows, widths=None):
    t=doc.add_table(rows=1, cols=len(headers)); t.style="Table Grid"; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,h in enumerate(headers):
        t.rows[0].cells[i].text=h; shade(t.rows[0].cells[i],"0F4C81")
        for r in t.rows[0].cells[i].paragraphs[0].runs: r.bold=True; r.font.color.rgb=RGBColor(255,255,255)
    for row in rows:
        cells=t.add_row().cells
        for i,v in enumerate(row): cells[i].text=str(v)
    return t


def build_doc(paths):
    doc=Document(); sec=doc.sections[0]
    sec.orientation=WD_ORIENT.LANDSCAPE; sec.page_width,sec.page_height=sec.page_height,sec.page_width
    sec.top_margin=Inches(.45); sec.bottom_margin=Inches(.45); sec.left_margin=Inches(.55); sec.right_margin=Inches(.55)
    normal=doc.styles["Normal"]; normal.font.name="Aptos"; normal.font.size=Pt(10); normal.paragraph_format.space_after=Pt(4)
    for name in ["Title","Heading 1","Heading 2"]:
        doc.styles[name].font.name="Aptos Display"; doc.styles[name].font.color.rgb=RGBColor(15,76,129)

    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run("NexCore OS"); r.bold=True; r.font.size=Pt(36); r.font.color.rgb=RGBColor(15,76,129)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run("Complete Orchestration Blueprint and Analysis Pack"); r.bold=True; r.font.size=Pt(24); r.font.color.rgb=RGBColor(71,85,105)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Ten architecture diagrams • responsibilities • invariants • risks • phased implementation").italic=True
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(paths[0]),width=Inches(9.8))
    doc.add_page_break()

    doc.add_heading("Executive architecture decision",1)
    doc.add_paragraph("NexCore OS should be built as a domain-specific execution intelligence operating system. The existing Next.js, NestJS, Temporal, NATS and Python foundation is directionally correct. The primary work is to stabilize ownership, contracts, lifecycle semantics, tenancy and extensibility—not to rewrite the stack.")
    add_table(doc,["Plane","Primary responsibility","Recommended technology","Source of truth"],[
        ("Experience","Author, operate and investigate","Next.js / React","No"),("Control","Resources, policy, API, audit","NestJS modular control plane","PostgreSQL"),("Orchestration","Durable lifecycle and DAG state","Temporal","Temporal + canonical DB projection"),("Execution","Runtime-specific automation","Python / isolated agents","Structured outcomes only"),("Intelligence","Diagnosis and repair proposals","Python / provider-neutral AI","No; derived recommendations"),("Data","Metadata, evidence, events and retrieval","PostgreSQL / MinIO / NATS / Qdrant","Store-dependent")])
    doc.add_heading("Non-negotiable invariants",2)
    for x in ["Only the orchestration kernel owns execution and node state transitions.","Workers report outcomes; they never update canonical lifecycle state directly.","AI proposes; policy and authorized humans decide whether mutation is allowed.","Workflow versions, attempts, events and evidence are immutable records.","PostgreSQL is canonical; MinIO is evidentiary; NATS and Qdrant are not sources of truth.","Every cross-plane contract is versioned and every command is idempotent."]:
        doc.add_paragraph(x,style="List Bullet")
    doc.add_page_break()

    for idx,((title,intro,decisions,risks),path) in enumerate(zip(SECTIONS,paths),1):
        doc.add_heading(title,1); doc.add_paragraph(intro)
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(path),width=Inches(10.35))
        add_table(doc,["Required decision / rule","Reason"],[(x,"Required for predictable ownership, security or recovery.") for x in decisions])
        doc.add_heading("Analysis risks",2)
        for r in risks: doc.add_paragraph(r,style="List Bullet")
        if idx < len(SECTIONS): doc.add_page_break()

    doc.add_page_break(); doc.add_heading("Implementation roadmap and analysis gates",1)
    add_table(doc,["Phase","Scope","Exit gate"],[
        ("0 — Baseline audit","Map current APIs, states, events, schemas and deployment dependencies.","Every current component has an owner and contract-gap register."),
        ("1 — Kernel contracts","Canonical resources, lifecycle, event envelope, idempotency, outbox, audit and tenant context.","One run is reconstructable from canonical state and append-only history."),
        ("2 — Web vertical slice","Define → execute → observe → diagnose → approve repair → replay.","Reliable complete loop without manual DB intervention."),
        ("3 — Runtime fabric","Capability registry, leases, heartbeats, remote agents and conformance suite.","A new conforming driver requires no kernel change."),
        ("4 — Security hardening","RLS, workload identity, secret manager, plugin permissions, retention and residency.","Threat model controls pass automated and manual validation."),
        ("5 — Ecosystem","SDK, signed packages, compatibility, CI/CD and integration connectors.","External team can safely build/install/remove a plugin."),
        ("6 — Governed autonomy","Evidence-grounded repair, policy gates, evaluation and rollback.","Autonomy improves measured recovery without reducing auditability.")])
    doc.add_heading("Measurements required before changing technology",2)
    for x in ["Control-plane p95/p99 latency and memory under realistic concurrency.","Temporal scheduling latency, workflow-history growth and replay reliability.","Agent utilization, queue depth, lease-loss frequency and recovery success.","Artifact volume, retention cost and evidence retrieval latency.","Diagnosis precision, citation validity, approval rate, repair success and rollback rate.","Developer lead time, defect escape rate and operational toil by subsystem."]:
        doc.add_paragraph(x,style="List Bullet")
    doc.add_heading("Practical conclusion",1)
    doc.add_paragraph("Do not begin with a .NET rewrite or a broad marketplace. First harden the kernel contracts and prove the web execution intelligence loop. If later measurements show the NestJS control plane is an organizational or operational constraint, ASP.NET Core can replace it behind stable OpenAPI and event contracts without changing the execution OS model.")
    footer=doc.sections[0].footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    rr=footer.add_run("NexCore OS Complete Orchestration Blueprint • Architecture baseline v1.0 • July 2026"); rr.font.size=Pt(8); rr.font.color.rgb=RGBColor(100,116,139)
    doc.core_properties.title="NexCore OS Complete Orchestration Blueprint"; doc.core_properties.author="NexCore"
    doc.save(DOCX)


if __name__ == "__main__":
    ps=diagrams(); build_doc(ps); print(DOCX)
    for p in ps: print(p)
