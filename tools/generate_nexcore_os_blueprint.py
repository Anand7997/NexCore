from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


OUT = Path.home() / "Downloads"
PNG = OUT / "NexCore OS Orchestration Blueprint.png"
DOCX = OUT / "NexCore OS Orchestration Blueprint.docx"

W, H = 2400, 1500
BG = "#06111f"
PANEL = "#0b1e33"
PANEL_2 = "#102a45"
CYAN = "#38d9ff"
BLUE = "#3787ff"
GREEN = "#42e6a4"
AMBER = "#ffc857"
PURPLE = "#b68cff"
RED = "#ff6b7a"
WHITE = "#eef7ff"
MUTED = "#9bb4ca"
GRID = "#10304d"


def font(size, bold=False):
    names = [
        "C:/Windows/Fonts/seguisb.ttf" if bold else "C:/Windows/Fonts/segoeui.ttf",
        "C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf",
    ]
    for name in names:
        if Path(name).exists():
            return ImageFont.truetype(name, size)
    return ImageFont.load_default()


F_TITLE = font(58, True)
F_SUB = font(24)
F_LANE = font(25, True)
F_BOX = font(23, True)
F_SMALL = font(18)
F_TINY = font(16)


def rounded(draw, box, fill=PANEL, outline=BLUE, radius=20, width=3):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def centered(draw, box, title, subtitle="", color=WHITE, accent=CYAN):
    x1, y1, x2, y2 = box
    cx = (x1 + x2) // 2
    title_box = draw.textbbox((0, 0), title, font=F_BOX)
    tw = title_box[2] - title_box[0]
    ty = y1 + 18
    draw.text((cx - tw / 2, ty), title, font=F_BOX, fill=color)
    if subtitle:
        lines = subtitle.split("\n")
        yy = ty + 38
        for line in lines:
            sb = draw.textbbox((0, 0), line, font=F_SMALL)
            sw = sb[2] - sb[0]
            draw.text((cx - sw / 2, yy), line, font=F_SMALL, fill=MUTED)
            yy += 25
    draw.line((x1 + 18, y1 + 8, x2 - 18, y1 + 8), fill=accent, width=5)


def arrow(draw, start, end, color=CYAN, width=4, label=None, dashed=False):
    x1, y1 = start
    x2, y2 = end
    if dashed:
        segments = 16
        for i in range(0, segments, 2):
            a = i / segments
            b = min((i + 1) / segments, 1)
            draw.line((x1 + (x2-x1)*a, y1 + (y2-y1)*a,
                       x1 + (x2-x1)*b, y1 + (y2-y1)*b), fill=color, width=width)
    else:
        draw.line((x1, y1, x2, y2), fill=color, width=width)
    import math
    angle = math.atan2(y2-y1, x2-x1)
    size = 14
    p1 = (x2, y2)
    p2 = (x2 - size*math.cos(angle-0.55), y2 - size*math.sin(angle-0.55))
    p3 = (x2 - size*math.cos(angle+0.55), y2 - size*math.sin(angle+0.55))
    draw.polygon([p1, p2, p3], fill=color)
    if label:
        mx, my = (x1+x2)//2, (y1+y2)//2
        bb = draw.textbbox((0, 0), label, font=F_TINY)
        pad = 6
        draw.rounded_rectangle((mx-(bb[2]-bb[0])//2-pad, my-22,
                                mx+(bb[2]-bb[0])//2+pad, my+4), 6, fill=BG)
        draw.text((mx-(bb[2]-bb[0])/2, my-20), label, font=F_TINY, fill=color)


def create_blueprint():
    img = Image.new("RGB", (W, H), BG)
    d = ImageDraw.Draw(img)
    for x in range(0, W, 50):
        d.line((x, 0, x, H), fill=GRID, width=1)
    for y in range(0, H, 50):
        d.line((0, y, W, y), fill=GRID, width=1)

    d.text((90, 55), "NEXCORE OS", font=F_TITLE, fill=WHITE)
    d.text((90, 125), "ORCHESTRATION BLUEPRINT  /  EXECUTION INTELLIGENCE CONTROL LOOP", font=F_SUB, fill=CYAN)
    d.text((1940, 75), "BLUEPRINT 01", font=F_LANE, fill=MUTED)
    d.text((1940, 112), "VERSION 1.0", font=F_SMALL, fill=MUTED)

    # Swimlane backgrounds
    lanes = [
        ("EXPERIENCE PLANE", 220, 430, BLUE),
        ("ORCHESTRATION KERNEL", 455, 800, CYAN),
        ("EXECUTION PLANE", 825, 1125, GREEN),
        ("EVIDENCE + INTELLIGENCE", 1150, 1405, PURPLE),
    ]
    for name, top, bottom, color in lanes:
        d.rounded_rectangle((55, top, W-55, bottom), 20, fill="#071827", outline=color, width=2)
        d.text((80, top+16), name, font=F_LANE, fill=color)

    # Experience
    ui = (155, 285, 500, 390)
    api = (650, 285, 1010, 390)
    policy = (1160, 285, 1510, 390)
    realtime = (1700, 285, 2180, 390)
    for box, title, sub, color in [
        (ui, "Next.js Workspace", "Author • Run • Observe", BLUE),
        (api, "Command/API Gateway", "REST • OpenAPI • Idempotency", CYAN),
        (policy, "Identity + Policy", "Tenant • RBAC • Quotas • Audit", AMBER),
        (realtime, "Realtime Projection", "Authorized execution stream", PURPLE),
    ]:
        rounded(d, box, outline=color)
        centered(d, box, title, sub, accent=color)
    arrow(d, (500, 338), (650, 338), BLUE, label="commands")
    arrow(d, (1010, 338), (1160, 338), AMBER, label="authorize")

    # Kernel
    validation = (120, 535, 410, 665)
    temporal = (535, 500, 975, 705)
    scheduler = (1100, 535, 1480, 665)
    nats = (1630, 535, 2200, 665)
    for box, title, sub, color in [
        (validation, "Validate + Persist", "Workflow version\nEnvironment • Input", BLUE),
        (temporal, "TEMPORAL KERNEL", "DAG • State • Retry • Timeout\nPause • Approval • Replay", CYAN),
        (scheduler, "Capability Scheduler", "Match • Lease • Heartbeat\nHealth • Capacity • Locality", GREEN),
        (nats, "NATS EVENT FABRIC", "Versioned facts • Outbox\nAt-least-once delivery", AMBER),
    ]:
        rounded(d, box, fill=PANEL_2 if box == temporal else PANEL, outline=color, width=4 if box == temporal else 3)
        centered(d, box, title, sub, accent=color)
    arrow(d, (825, 390), (265, 535), BLUE, label="start execution")
    arrow(d, (410, 600), (535, 600), CYAN, label="durable start")
    arrow(d, (975, 600), (1100, 600), GREEN, label="place node")
    arrow(d, (975, 675), (1630, 625), AMBER, label="state events")
    arrow(d, (1915, 535), (1940, 390), PURPLE, label="project")

    # Execution plane
    workers = [
        ((100, 925, 500, 1055), "Web Runtime Pool", "Playwright • Browser grid", GREEN),
        ((570, 925, 970, 1055), "API + Data Pool", "HTTP • Database • Files", BLUE),
        ((1040, 925, 1440, 1055), "Mobile Device Pool", "Appium • Android • iOS", AMBER),
        ((1510, 925, 1910, 1055), "Desktop Agent Pool", "Windows UIA • Customer network", PURPLE),
    ]
    for box, title, sub, color in workers:
        rounded(d, box, outline=color)
        centered(d, box, title, sub, accent=color)
    agent = (1990, 925, 2280, 1055)
    rounded(d, agent, outline=RED)
    centered(d, agent, "Agent Contract", "Envelope • Result\nLease • Cancel", accent=RED)
    for x in [300, 770, 1240, 1710, 2135]:
        arrow(d, (1290, 665), (x, 925), GREEN, label="lease" if x == 1240 else None)
        arrow(d, (x+22, 925), (1700, 665), AMBER, dashed=True)

    # Evidence and intelligence
    pg = (110, 1235, 430, 1355)
    minio = (520, 1235, 840, 1355)
    qdrant = (930, 1235, 1250, 1355)
    ai = (1340, 1200, 1730, 1370)
    approval = (1820, 1200, 2260, 1370)
    for box, title, sub, color in [
        (pg, "PostgreSQL", "Canonical state + audit", BLUE),
        (minio, "MinIO", "Screenshots • Traces • Video", GREEN),
        (qdrant, "Qdrant", "Derived retrieval memory", PURPLE),
        (ai, "AI Diagnosis", "Correlate evidence\nCite • Score • Propose repair", PURPLE),
        (approval, "Policy + Human Gate", "Approve • Reject • Constrain\nCreate new version + replay", AMBER),
    ]:
        rounded(d, box, outline=color)
        centered(d, box, title, sub, accent=color)
    arrow(d, (300, 1055), (680, 1235), GREEN, label="artifacts")
    arrow(d, (770, 1055), (270, 1235), BLUE, label="results")
    arrow(d, (1915, 665), (1535, 1200), PURPLE, label="failure event")
    arrow(d, (840, 1295), (1340, 1295), PURPLE, label="evidence")
    arrow(d, (1250, 1295), (1340, 1295), PURPLE)
    arrow(d, (1730, 1285), (1820, 1285), AMBER, label="proposal")
    arrow(d, (2040, 1200), (755, 705), CYAN, label="approved repair / replay")

    # Legend / invariant footer
    d.text((90, 1440), "KERNEL INVARIANT:", font=F_SMALL, fill=CYAN)
    d.text((285, 1440), "Only Temporal owns execution state transitions. Workers report outcomes; AI proposes; policy authorizes.", font=F_SMALL, fill=WHITE)
    d.text((1910, 1440), "NEXCORE OS  /  2026", font=F_SMALL, fill=MUTED)
    img.save(PNG, quality=95)


def create_docx():
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("NexCore OS Orchestration Blueprint")
    r.bold = True
    r.font.name = "Aptos Display"
    r.font.size = Pt(24)
    r.font.color.rgb = RGBColor(15, 76, 129)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Execution intelligence control loop • Architecture baseline v1.0").italic = True
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(PNG), width=Inches(10.6))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Invariant: only the orchestration kernel owns execution state transitions. Runtime workers report outcomes; AI proposes changes; policy and humans authorize them.")
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(71, 85, 105)
    doc.core_properties.title = "NexCore OS Orchestration Blueprint"
    doc.core_properties.author = "NexCore"
    doc.save(DOCX)


if __name__ == "__main__":
    create_blueprint()
    create_docx()
    print(PNG)
    print(DOCX)
