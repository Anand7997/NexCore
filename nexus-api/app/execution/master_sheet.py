"""Master-sheet helpers for data-driven desktop execution."""
from __future__ import annotations

import csv
import json
import urllib.request
from copy import deepcopy
from dataclasses import dataclass
from io import StringIO
from pathlib import Path
from typing import Any
from urllib.parse import parse_qs, urlencode, urlparse


class MasterSheetError(ValueError):
    """Raised when a configured master sheet cannot be read or parsed."""


_SECTION_ALIASES = {
    "app": "applications",
    "apps": "applications",
    "application": "applications",
    "applications": "applications",
    "window": "windows",
    "windows": "windows",
    "element": "elements",
    "elements": "elements",
    "object": "elements",
    "objects": "elements",
    "desktop_element": "elements",
    "desktop_elements": "elements",
    "locator": "locators",
    "locators": "locators",
    "path": "paths",
    "paths": "paths",
    "file_path": "paths",
    "file_paths": "paths",
    "data": "test_data",
    "testdata": "test_data",
    "test_data": "test_data",
    "environment": "environments",
    "environments": "environments",
}

VALID_DESKTOP_LOCATOR_STRATEGIES = {
    "accessibility id",
    "automation id",
    "name",
    "xpath",
    "uia path",
    "class name",
    "ocr",
    "visual",
    "image",
    "text",
}

MASTER_SHEET_TEMPLATE: dict[str, Any] = {
    "applications": {
        "invoice_app": {
            "key": "invoice_app",
            "name": "Invoice Desktop App",
            "application_path": r"C:\Apps\Invoice\Invoice.exe",
            "args": "",
            "working_directory": r"C:\Apps\Invoice",
            "environment": "qa",
        }
    },
    "windows": {
        "invoice_main": {
            "key": "invoice_main",
            "application_key": "invoice_app",
            "title": "Invoice Manager",
            "process_name": "Invoice.exe",
        }
    },
    "elements": {
        "customer_name_input": {
            "key": "customer_name_input",
            "name": "Customer Name",
            "application_key": "invoice_app",
            "window_key": "invoice_main",
            "control_type": "Edit",
            "automation_id": "txtCustomerName",
            "locator_strategy": "accessibility id",
            "primary_locator_value": "txtCustomerName",
            "alternative_locators": [
                {"strategy": "name", "locator": "Customer Name"},
                {"strategy": "class name", "locator": "Edit"},
            ],
            "active": True,
        },
        "submit_invoice_button": {
            "key": "submit_invoice_button",
            "name": "Submit Invoice",
            "control_type": "Button",
            "automation_id": "btnSubmit",
            "locator_strategy": "accessibility id",
            "primary_locator_value": "btnSubmit",
            "active": True,
        },
    },
    "paths": {
        "invoice_export_folder": {
            "key": "invoice_export_folder",
            "path": r"C:\Exports\Invoices",
            "environment": "qa",
        }
    },
    "test_data": {
        "customers": {
            "default": {
                "name": "Asha Rao",
                "balance": "1000.00",
            }
        }
    },
    "environments": {
        "qa": {
            "machine_group": "windows-qa",
            "runtime_agent_tag": "desktop-uia3",
        }
    },
}


def _normalise_key(value: Any) -> str:
    return str(value or "").strip()


def _normalise_section(value: Any) -> str:
    key = _normalise_key(value).lower().replace(" ", "_").replace("-", "_")
    return _SECTION_ALIASES.get(key, key)


def _item_key(item: dict[str, Any]) -> str:
    return _normalise_key(
        item.get("key")
        or item.get("object_key")
        or item.get("element_key")
        or item.get("app_key")
        or item.get("application_key")
        or item.get("path_key")
        or item.get("data_key")
        or item.get("name")
    )


def _coerce_cell(value: Any) -> Any:
    if not isinstance(value, str):
        return value
    text = value.strip()
    if text == "":
        return ""
    if text[:1] in {"[", "{"}:
        try:
            return json.loads(text)
        except json.JSONDecodeError:
            return value
    if "|" in text:
        return [part.strip() for part in text.split("|") if part.strip()]
    return value


def _deep_merge(base: dict[str, Any], override: dict[str, Any]) -> dict[str, Any]:
    merged = deepcopy(base)
    for key, value in (override or {}).items():
        section = _normalise_section(key)
        if isinstance(value, dict) and isinstance(merged.get(section), dict):
            merged[section] = _deep_merge(merged[section], value)
        else:
            merged[section] = deepcopy(value)
    return merged


def _load_json(path: Path) -> dict[str, Any]:
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except OSError as exc:
        raise MasterSheetError(f"Cannot read master sheet: {path}") from exc
    except json.JSONDecodeError as exc:
        raise MasterSheetError(f"Invalid JSON master sheet: {path}") from exc
    if not isinstance(data, dict):
        raise MasterSheetError("JSON master sheet must contain an object at the root")
    return data


def _rows_to_section_mapping(rows: list[dict[str, Any]]) -> dict[str, Any]:
    data: dict[str, Any] = {}
    for row in rows:
        section = _normalise_section(row.get("section") or row.get("type") or row.get("sheet") or "elements")
        key = _item_key(row)
        if not key:
            continue
        item = {
            column: _coerce_cell(value)
            for column, value in row.items()
            if column and value not in (None, "") and column.lower() not in {"section", "type", "sheet"}
        }
        item.setdefault("key", key)
        data.setdefault(section, {})[key] = item
    return data


def _load_csv_text(text: str) -> dict[str, Any]:
    rows = list(csv.DictReader(StringIO(text)))
    return _rows_to_section_mapping(rows)


def _load_csv(path: Path) -> dict[str, Any]:
    try:
        text = path.read_text(encoding="utf-8-sig")
    except OSError as exc:
        raise MasterSheetError(f"Cannot read master sheet: {path}") from exc
    return _load_csv_text(text)


def _load_xlsx(path: Path) -> dict[str, Any]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise MasterSheetError("XLSX master sheets require openpyxl to be installed") from exc
    try:
        workbook = load_workbook(path, data_only=True, read_only=True)
    except OSError as exc:
        raise MasterSheetError(f"Cannot read master sheet: {path}") from exc

    data: dict[str, Any] = {}
    try:
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            headers = [_normalise_key(value).lower() for value in next(rows, [])]
            if not headers:
                continue
            section = _normalise_section(sheet.title)
            for row_values in rows:
                row = {
                    headers[index]: _coerce_cell(value)
                    for index, value in enumerate(row_values)
                    if index < len(headers) and headers[index] and value not in (None, "")
                }
                key = _item_key(row)
                if not key:
                    continue
                row.setdefault("key", key)
                data.setdefault(section, {})[key] = row
    finally:
        close = getattr(workbook, "close", None)
        if callable(close):
            close()
    return data


def _load_yaml(path: Path) -> dict[str, Any]:
    try:
        import yaml
    except ImportError as exc:
        raise MasterSheetError("YAML master sheets require PyYAML to be installed") from exc
    try:
        data = yaml.safe_load(path.read_text(encoding="utf-8"))
    except OSError as exc:
        raise MasterSheetError(f"Cannot read master sheet: {path}") from exc
    except yaml.YAMLError as exc:
        raise MasterSheetError(f"Invalid YAML master sheet: {path}") from exc
    if not isinstance(data, dict):
        raise MasterSheetError("YAML master sheet must contain a mapping at the root")
    return data


def _google_sheet_export_url(url: str) -> str:
    parsed = urlparse(url)
    if "docs.google.com" not in parsed.netloc or "/spreadsheets/" not in parsed.path:
        return url
    if "/export" in parsed.path:
        return url
    parts = parsed.path.split("/")
    try:
        spreadsheet_id = parts[parts.index("d") + 1]
    except (ValueError, IndexError):
        return url
    query = parse_qs(parsed.query)
    gid = query.get("gid", ["0"])[0]
    return f"https://docs.google.com/spreadsheets/d/{spreadsheet_id}/export?{urlencode({'format': 'csv', 'gid': gid})}"


def load_master_sheet_url(url: str, headers: dict[str, str] | None = None) -> dict[str, Any]:
    """Load a JSON/CSV master sheet from HTTP(S), including Google Sheets CSV export URLs."""
    normalized_url = _google_sheet_export_url(str(url or "").strip())
    if not normalized_url:
        raise MasterSheetError("master_sheet_url is empty")
    request = urllib.request.Request(normalized_url, headers=headers or {"Accept": "application/json,text/csv,*/*"})
    try:
        with urllib.request.urlopen(request, timeout=30) as response:
            raw = response.read()
            content_type = str(response.headers.get("content-type") or "").lower()
    except OSError as exc:
        raise MasterSheetError(f"Cannot fetch remote master sheet: {normalized_url}") from exc
    text = raw.decode("utf-8-sig", errors="replace")
    if "json" in content_type or normalized_url.lower().endswith(".json"):
        try:
            data = json.loads(text)
        except json.JSONDecodeError as exc:
            raise MasterSheetError(f"Invalid JSON remote master sheet: {normalized_url}") from exc
        if not isinstance(data, dict):
            raise MasterSheetError("Remote JSON master sheet must contain an object at the root")
        return data
    return _load_csv_text(text)


def load_master_sheet_sql(connection_string: str, query: str) -> dict[str, Any]:
    """Load master-sheet rows from a SELECT query with section/key columns."""
    sql = str(query or "").strip()
    if not sql.upper().startswith("SELECT"):
        raise MasterSheetError("Master-sheet database query must be a SELECT")
    try:
        from sqlalchemy import create_engine, text
    except ImportError as exc:
        raise MasterSheetError("Database-backed master sheets require SQLAlchemy") from exc
    engine = None
    try:
        engine = create_engine(connection_string)
        with engine.connect() as connection:
            rows = [dict(row._mapping) for row in connection.execute(text(sql))]
    except Exception as exc:
        raise MasterSheetError("Cannot load database-backed master sheet") from exc
    finally:
        try:
            if engine is not None:
                engine.dispose()
        except Exception:
            pass
    return _rows_to_section_mapping(rows)


def load_master_sheet(path_value: str) -> dict[str, Any]:
    path = Path(path_value).expanduser()
    suffix = path.suffix.lower()
    if suffix == ".json":
        return _load_json(path)
    if suffix == ".csv":
        return _load_csv(path)
    if suffix in {".xlsx", ".xlsm"}:
        return _load_xlsx(path)
    if suffix in {".yaml", ".yml"}:
        return _load_yaml(path)
    raise MasterSheetError(f"Unsupported master sheet format: {path.suffix or path}")


def _section_as_mapping(value: Any) -> dict[str, Any]:
    if isinstance(value, dict):
        return value
    if isinstance(value, list):
        mapped: dict[str, Any] = {}
        for item in value:
            if not isinstance(item, dict):
                continue
            key = _item_key(item)
            if key:
                mapped[key] = item
        return mapped
    return {}


def normalise_master_sheet(data: dict[str, Any] | None) -> dict[str, Any]:
    """Return a stable section/key mapping regardless of CSV/XLSX/JSON shape."""

    normalized: dict[str, Any] = {}
    for raw_section, value in (data or {}).items():
        section = _normalise_section(raw_section)
        if section in {"applications", "windows", "elements", "locators", "paths", "test_data", "environments"}:
            normalized[section] = _section_as_mapping(value)
        else:
            normalized[section] = deepcopy(value)
    return normalized


def master_sheet_summary(data: dict[str, Any] | None) -> dict[str, int]:
    normalized = normalise_master_sheet(data)
    return {
        "applications": len(_section_as_mapping(normalized.get("applications"))),
        "windows": len(_section_as_mapping(normalized.get("windows"))),
        "elements": len(_section_as_mapping(normalized.get("elements"))),
        "locators": len(_section_as_mapping(normalized.get("locators"))),
        "paths": len(_section_as_mapping(normalized.get("paths"))),
        "test_data": len(_section_as_mapping(normalized.get("test_data"))),
        "environments": len(_section_as_mapping(normalized.get("environments"))),
    }


def _issue(severity: str, section: str, key: str, message: str, field: str = "") -> dict[str, str]:
    return {
        "severity": severity,
        "section": section,
        "key": key,
        "field": field,
        "message": message,
    }


def _active_false(value: Any) -> bool:
    return str(value).strip().lower() in {"false", "0", "no", "inactive", "disabled"}


def validate_master_sheet(data: dict[str, Any] | None) -> list[dict[str, str]]:
    """Validate master-sheet data before runtime resolution."""

    raw = data or {}
    normalized = normalise_master_sheet(raw)
    sheet = DesktopMasterSheet(normalized)
    issues: list[dict[str, str]] = []

    for raw_section, value in raw.items():
        section = _normalise_section(raw_section)
        if not isinstance(value, list):
            continue
        seen: set[str] = set()
        for item in value:
            if not isinstance(item, dict):
                continue
            key = _item_key(item)
            if not key:
                continue
            if key in seen:
                issues.append(_issue("error", section, key, "Duplicate key in master-sheet section", "key"))
            seen.add(key)

    for app_key, app in sheet.section("applications").items():
        if not isinstance(app, dict):
            continue
        if not sheet.application_path(app_key):
            issues.append(_issue("error", "applications", app_key, "Application is missing executable/application path", "application_path"))
        if _active_false(app.get("active") or app.get("status")):
            issues.append(_issue("warning", "applications", app_key, "Application is marked inactive", "active"))

    for object_key, element in sheet.section("elements").items():
        if not isinstance(element, dict):
            continue
        locators = sheet.locators_for(element)
        if not locators:
            issues.append(_issue("error", "elements", object_key, "Desktop object has no usable locator", "locator"))
        if not _normalise_key(element.get("control_type") or element.get("element_type")):
            issues.append(_issue("warning", "elements", object_key, "Desktop object is missing control type", "control_type"))
        if _active_false(element.get("active") or element.get("status")):
            issues.append(_issue("warning", "elements", object_key, "Desktop object is marked inactive", "active"))
        for locator in locators:
            strategy = _normalise_key(locator.get("strategy")).lower().replace("_", " ")
            if strategy and strategy not in VALID_DESKTOP_LOCATOR_STRATEGIES:
                issues.append(_issue("warning", "elements", object_key, f"Unknown locator strategy: {locator.get('strategy')}", "locator_strategy"))

    for path_key, path_item in sheet.section("paths").items():
        if isinstance(path_item, dict) and not sheet.path_value(path_key):
            issues.append(_issue("error", "paths", path_key, "Path entry is missing path/value", "path"))
        elif not isinstance(path_item, dict) and not _normalise_key(path_item):
            issues.append(_issue("error", "paths", path_key, "Path entry is empty", "path"))

    for locator_key, locator in sheet.section("locators").items():
        if not isinstance(locator, dict):
            continue
        strategy = _normalise_key(locator.get("strategy")).lower().replace("_", " ")
        locator_value = _normalise_key(locator.get("locator") or locator.get("selector") or locator.get("value"))
        if not locator_value:
            issues.append(_issue("error", "locators", locator_key, "Locator entry is missing locator value", "locator"))
        if strategy and strategy not in VALID_DESKTOP_LOCATOR_STRATEGIES:
            issues.append(_issue("warning", "locators", locator_key, f"Unknown locator strategy: {locator.get('strategy')}", "strategy"))

    # Inactive element objects used by workflows
    for object_key, element in sheet.section("elements").items():
        if not isinstance(element, dict):
            continue
        if _active_false(element.get("active") or element.get("status")):
            issues.append(_issue("warning", "elements", object_key, "Desktop object is marked inactive and may affect execution", "active"))

    # Environments missing machine_group / agent_group
    for env_key, env in sheet.section("environments").items():
        if not isinstance(env, dict):
            continue
        if not _normalise_key(env.get("machine_group") or env.get("agent_group")):
            issues.append(_issue("warning", "environments", env_key, "Environment has no machine_group or agent_group defined", "machine_group"))

    # Broken screenshot/baseline file references
    for object_key, element in sheet.section("elements").items():
        if not isinstance(element, dict):
            continue
        screenshot_ref = _normalise_key(
            element.get("screenshot_reference")
            or element.get("screenshot_path")
            or ""
        )
        if screenshot_ref and not Path(screenshot_ref).expanduser().exists():
            issues.append(_issue("warning", "elements", object_key, f"Screenshot reference not found: {screenshot_ref}", "screenshot_reference"))

    return issues


def _walk_dotted(data: Any, path: str) -> Any:
    current = data
    for part in [segment for segment in path.split(".") if segment]:
        if isinstance(current, dict) and part in current:
            current = current[part]
            continue
        if isinstance(current, list):
            try:
                current = current[int(part)]
                continue
            except (ValueError, IndexError):
                return None
        return None
    return current


def _first_value(data: dict[str, Any], keys: tuple[str, ...]) -> Any:
    for key in keys:
        value = data.get(key)
        if value not in (None, ""):
            return value
    return None


@dataclass
class DesktopMasterSheet:
    """Resolved master-sheet data used by desktop workflow compilation."""

    data: dict[str, Any]
    source: str = "inline"

    @classmethod
    def from_variables(cls, variables: dict[str, Any] | None) -> "DesktopMasterSheet":
        variables = variables or {}
        data: dict[str, Any] = {}
        source = "inline"
        path = variables.get("master_sheet_path") or variables.get("masterSheetPath")
        if path:
            data = _deep_merge(data, load_master_sheet(str(path)))
            source = str(path)
        url = variables.get("master_sheet_url") or variables.get("masterSheetUrl")
        if url:
            data = _deep_merge(data, load_master_sheet_url(str(url)))
            source = str(url)
        sql_conn = variables.get("master_sheet_db") or variables.get("masterSheetDb")
        sql_query = variables.get("master_sheet_query") or variables.get("masterSheetQuery")
        if sql_conn and sql_query:
            data = _deep_merge(data, load_master_sheet_sql(str(sql_conn), str(sql_query)))
            source = "database"
        for key in ("master_sheet", "masterSheet", "desktop_master_sheet", "desktopMasterSheet"):
            value = variables.get(key)
            if isinstance(value, dict):
                data = _deep_merge(data, value)
                source = "inline"
        return cls(data=data, source=source)

    @property
    def is_empty(self) -> bool:
        return not bool(self.data)

    def section(self, name: str) -> dict[str, Any]:
        return _section_as_mapping(self.data.get(_normalise_section(name)))

    def application(self, key: str | None) -> dict[str, Any] | None:
        key = _normalise_key(key)
        if not key:
            return None
        value = self.section("applications").get(key)
        if isinstance(value, dict):
            return value
        if isinstance(value, str):
            return {"key": key, "application_path": value}
        return None

    def application_path(self, key: str | None) -> str:
        app = self.application(key)
        if not app:
            return ""
        return str(_first_value(app, ("application_path", "executable_path", "path", "app", "exe", "file_path")) or "")

    def element(self, key: str | None) -> dict[str, Any] | None:
        key = _normalise_key(key)
        if not key:
            return None
        value = self.section("elements").get(key)
        return value if isinstance(value, dict) else None

    def has_element(self, key: str | None) -> bool:
        return self.element(key) is not None

    def path_value(self, key: str | None) -> str:
        key = _normalise_key(key)
        if not key:
            return ""
        value = self.section("paths").get(key)
        if isinstance(value, dict):
            return str(_first_value(value, ("value", "path", "file_path", "directory", "folder", "location")) or "")
        return str(value or "")

    def data_value(self, key: str | None) -> Any:
        key = _normalise_key(key)
        if not key:
            return None
        data = self.data.get("test_data") or self.data.get("data") or {}
        if isinstance(data, dict):
            if key in data:
                value = data[key]
                if isinstance(value, dict):
                    return _first_value(value, ("value", "input_value", "sample_value", "text", "data")) or value
                return value
            return _walk_dotted(data, key)
        return None

    def locators_for(self, element: dict[str, Any]) -> list[dict[str, Any]]:
        locators: list[dict[str, Any]] = []

        def add(strategy: Any, locator: Any, reason: str = "master_sheet") -> None:
            value = _normalise_key(locator)
            if not value:
                return
            locators.append({
                "strategy": str(strategy or ""),
                "locator": value,
                "reason": reason,
            })

        for item in element.get("locators") or element.get("alternative_locators") or []:
            if isinstance(item, dict):
                add(item.get("strategy"), item.get("locator") or item.get("selector") or item.get("value"), item.get("reason") or "master_sheet")
        add(
            element.get("primary_locator_strategy") or element.get("locator_strategy"),
            element.get("primary_locator_value") or element.get("locator") or element.get("selector"),
        )
        add("accessibility id", element.get("automation_id") or element.get("auto_id"))
        add("xpath", element.get("uia_path") or element.get("xpath"))
        add("name", element.get("name") or element.get("text"))
        add("class name", element.get("class_name") or element.get("class"))
        return locators
